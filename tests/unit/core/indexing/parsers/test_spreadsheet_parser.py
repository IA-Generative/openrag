"""Spreadsheet parser: .xlsx/.xlsm/.xls/.ods read natively (mycollections .xlsx import)."""

from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from core.chunking.recursive import RecursiveSplitter
from core.indexing.parsers.spreadsheet_parser import SpreadsheetParser, column_letter
from core.indexing.validators import validate_ooxml_package
from core.models.chunk import ChunkType
from core.models.document import Document, DocumentType
from core.utils.exceptions import ValidationError

FIXTURES = Path(__file__).parent / "fixtures"


def _xlsx(tmp_path: Path, build, name: str = "classeur.xlsx") -> Document:
    wb = openpyxl.Workbook()
    build(wb)
    path = tmp_path / name
    wb.save(path)
    return Document(id="d1", filename=name, content_type=DocumentType.SPREADSHEET, source_path=str(path))


def _budget_sheets(wb) -> None:
    ws = wb.active
    ws.title = "Crédits 2025"
    ws.append(["Département", "Programme", "AE (€)", "Date"])
    ws.append(["Finistère", "BOP 112", 1500.5, dt.date(2025, 3, 1)])
    ws.append(["Morbihan", None, 2000, dt.date(2025, 4, 1)])
    ws.merge_cells("B2:B3")
    synth = wb.create_sheet("Synthèse")
    synth.append(["Indicateur", "Valeur"])
    synth.append(["Total AE", "=SUM('Crédits 2025'!C2:C3)"])  # never recalculated: no cached value
    hidden = wb.create_sheet("Paramètres")
    hidden.append(["clé", "valeur"])
    hidden.append(["taux", 0.2])
    hidden.sheet_state = "hidden"


def test_extensions_map_to_the_spreadsheet_type():
    for name in ("a.xlsx", "a.xlsm", "a.xls", "a.ods"):
        assert Document.detect_content_type(name) is DocumentType.SPREADSHEET


@pytest.mark.asyncio
async def test_each_visible_sheet_becomes_a_page_with_its_table(tmp_path):
    processed = await SpreadsheetParser().parse(_xlsx(tmp_path, _budget_sheets))

    assert [b.page_number for b in processed.text_blocks] == [1, 2]
    first = processed.text_blocks[0].text
    assert first.startswith("## Crédits 2025\n\n| Département | Programme | AE (€) | Date |")
    assert "| Finistère | BOP 112 | 1500.5 | 2025-03-01 |" in first
    assert processed.text_blocks[1].text.startswith("## Synthèse")
    assert processed.page_count == 2


@pytest.mark.asyncio
async def test_merged_cells_take_their_top_left_value(tmp_path):
    processed = await SpreadsheetParser().parse(_xlsx(tmp_path, _budget_sheets))

    assert "| Morbihan | BOP 112 | 2000 | 2025-04-01 |" in processed.text_blocks[0].text
    assert processed.parse_info["spreadsheet"]["merged_ranges"] == 1


@pytest.mark.asyncio
async def test_hidden_sheets_and_formulas_without_value_are_reported(tmp_path):
    info = (await SpreadsheetParser().parse(_xlsx(tmp_path, _budget_sheets))).parse_info["spreadsheet"]

    assert info["sheets"] == 3
    assert info["sheets_indexed"] == 2
    assert info["hidden_sheets"] == 1 and info["hidden_sheets_skipped"] == 1
    assert info["formulas_without_value"] == 1
    assert info["truncated"] is False


@pytest.mark.asyncio
async def test_hidden_sheets_can_be_included(tmp_path):
    processed = await SpreadsheetParser(include_hidden_sheets=True).parse(_xlsx(tmp_path, _budget_sheets))

    assert [b.text.splitlines()[0] for b in processed.text_blocks] == [
        "## Crédits 2025",
        "## Synthèse",
        "## Paramètres",
    ]
    assert processed.parse_info["spreadsheet"]["hidden_sheets_skipped"] == 0


@pytest.mark.asyncio
async def test_a_large_sheet_is_truncated_and_says_so(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "Données"
        ws.append(["Id", "Commune", "Montant", "Année", "Statut"])
        for i in range(20_000):
            ws.append([i, f"Commune {i}", i * 10, 2025, "ok"])

    processed = await SpreadsheetParser(max_cells=50_000).parse(_xlsx(tmp_path, build))
    info = processed.parse_info["spreadsheet"]

    assert info["truncated"] is True
    assert info["rows_total"] == 20_001
    assert info["rows_kept"] == 10_000  # 50 000 cells / 5 columns, header included
    assert "| 9998 | Commune 9998 |" in processed.text_blocks[0].text
    assert "Commune 9999 " not in processed.text_blocks[0].text


@pytest.mark.asyncio
async def test_a_wide_sheet_is_split_into_column_groups_repeating_the_first_column(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "Large"
        ws.append(["Code"] + [f"M{i}" for i in range(1, 26)])
        ws.append(["UO-35"] + list(range(1, 26)))

    text = (await SpreadsheetParser(columns_per_group=12).parse(_xlsx(tmp_path, build))).text_blocks[0].text
    headings = re.findall(r"^## .*$", text, flags=re.M)

    assert headings == ["## Large (colonnes B–L)", "## Large (colonnes M–W)", "## Large (colonnes X–Z)"]
    assert text.count("| Code |") == 3
    assert text.count("| UO-35 |") == 3


@pytest.mark.asyncio
async def test_cells_are_made_safe_for_a_markdown_table(tmp_path):
    def build(wb):
        ws = wb.active
        ws.append(["Libellé", "Note"])
        ws.append(["a|b", "ligne 1\nligne 2"])
        ws.append([None, None])  # empty row skipped
        ws.append([True, 3.0])

    text = (await SpreadsheetParser().parse(_xlsx(tmp_path, build))).text_blocks[0].text

    assert "| a\\|b | ligne 1 ligne 2 |" in text
    assert "| VRAI | 3 |" in text
    assert "|  |  |" not in text


@pytest.mark.asyncio
async def test_empty_header_cells_get_their_column_letter(tmp_path):
    def build(wb):
        ws = wb.active
        ws.append(["Nom", None, "Total"])
        ws.append(["x", 1, 2])

    text = (await SpreadsheetParser().parse(_xlsx(tmp_path, build))).text_blocks[0].text
    assert "| Nom | B | Total |" in text


@pytest.mark.asyncio
async def test_xlsm_is_read_like_xlsx(tmp_path):
    processed = await SpreadsheetParser().parse(_xlsx(tmp_path, _budget_sheets, name="macros.xlsm"))
    assert processed.text_blocks[0].text.startswith("## Crédits 2025")


@pytest.mark.asyncio
async def test_legacy_xls_with_merged_cells_and_a_hidden_sheet():
    doc = Document(
        id="x", filename="budget.xls", content_type=DocumentType.SPREADSHEET, source_path=str(FIXTURES / "budget.xls")
    )

    processed = await SpreadsheetParser().parse(doc)
    info = processed.parse_info["spreadsheet"]

    assert len(processed.text_blocks) == 1
    text = processed.text_blocks[0].text
    assert "| Morbihan | BOP 112 | 2000 | 2024-04-01 |" in text
    assert info["merged_ranges"] == 1 and info["hidden_sheets_skipped"] == 1


@pytest.mark.asyncio
async def test_ods(tmp_path):
    path = tmp_path / "tableau.ods"
    with pd.ExcelWriter(path, engine="odf") as writer:
        pd.DataFrame({"Commune": ["Rennes", "Brest"], "Population": [222485, 139926]}).to_excel(
            writer, sheet_name="Communes", index=False
        )
        pd.DataFrame({"Clé": ["a"], "Valeur": [1]}).to_excel(writer, sheet_name="Autre", index=False)
    doc = Document(id="o", filename="tableau.ods", content_type=DocumentType.SPREADSHEET, source_path=str(path))

    processed = await SpreadsheetParser().parse(doc)

    assert [b.text.splitlines()[0] for b in processed.text_blocks] == ["## Communes", "## Autre"]
    assert "| Rennes | 222485 |" in processed.text_blocks[0].text


@pytest.mark.asyncio
async def test_bytes_only_documents_are_read_through_a_temporary_file(tmp_path):
    doc = _xlsx(tmp_path, _budget_sheets)
    raw = Path(doc.source_path).read_bytes()

    processed = await SpreadsheetParser().parse(
        Document(id="b", filename="classeur.xlsx", content_type=DocumentType.SPREADSHEET, raw_bytes=raw)
    )
    assert len(processed.text_blocks) == 2


@pytest.mark.asyncio
async def test_chunks_repeat_the_header_row(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "Programmation"
        ws.append(["Opération", "Commune", "Montant"])
        for i in range(300):
            ws.append([f"Opération numéro {i}", f"Commune {i}", i * 1000])

    processed = await SpreadsheetParser().parse(_xlsx(tmp_path, build))
    chunks = RecursiveSplitter(chunk_size=200, chunk_overlap_rate=0.0, length_function=lambda t: len(t.split())).chunk(
        processed, partition="p"
    )
    tables = [c for c in chunks if c.chunk_type == ChunkType.TABLE]

    assert len(tables) > 3
    assert all(c.text.lstrip().startswith("| Opération | Commune | Montant |") for c in tables)
    assert all(c.page_number == 1 for c in tables)


def test_column_letters():
    assert [column_letter(i) for i in (0, 25, 26, 27, 701, 702)] == ["A", "Z", "AA", "AB", "ZZ", "AAA"]


def test_validator_accepts_a_real_workbook_and_refuses_a_plain_zip(tmp_path):
    doc = _xlsx(tmp_path, _budget_sheets)
    with open(doc.source_path, "rb") as fh:
        validate_ooxml_package("xlsx", fh)

    fake = io.BytesIO()
    with zipfile.ZipFile(fake, "w") as zf:
        zf.writestr("notes.txt", "pas un classeur")
    fake.seek(0)
    with pytest.raises(ValidationError):
        validate_ooxml_package("xlsx", fake)
